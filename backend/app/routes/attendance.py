import io
import json
import aiosqlite
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from fastapi import APIRouter, Response, HTTPException
from typing import Optional
from ..config import DB_PATH

router = APIRouter(prefix="/api/attendance", tags=["Attendance & Reports"])

@router.get("/summary")
async def get_attendance_summary(date: Optional[str] = None):
    """Returns today's high-level attendance metrics for dashboard cards."""
    target_date = date or datetime.now().strftime("%Y-%m-%d")

    async with aiosqlite.connect(DB_PATH) as db:
        # Total Enrolled Active Users
        async with db.execute("SELECT COUNT(*) FROM users WHERE is_active = 1;") as cursor:
            total_enrolled = (await cursor.fetchone())[0]

        # Present Today
        async with db.execute("""
            SELECT COUNT(DISTINCT user_id) FROM attendance_events 
            WHERE date = ?;
        """, (target_date,)) as cursor:
            present_count = (await cursor.fetchone())[0]

        # Late Today
        async with db.execute("""
            SELECT COUNT(DISTINCT user_id) FROM attendance_events 
            WHERE date = ? AND status = 'Late';
        """, (target_date,)) as cursor:
            late_count = (await cursor.fetchone())[0]

        # Unknown Visitors Today
        async with db.execute("""
            SELECT COUNT(*) FROM unknown_visitors 
            WHERE date = ?;
        """, (target_date,)) as cursor:
            unknown_count = (await cursor.fetchone())[0]

        absent_count = max(0, total_enrolled - present_count)
        percentage = round((present_count / total_enrolled * 100), 1) if total_enrolled > 0 else 0.0

        return {
            "date": target_date,
            "total_enrolled": total_enrolled,
            "present_count": present_count,
            "absent_count": absent_count,
            "late_count": late_count,
            "unknown_count": unknown_count,
            "attendance_percentage": percentage
        }

@router.get("/present")
async def list_present_records(date: Optional[str] = None, department: Optional[str] = None, role: Optional[str] = None):
    """Returns list of students/faculty who checked in on target date with full footprint audit details."""
    target_date = date or datetime.now().strftime("%Y-%m-%d")
    
    query = """
        SELECT a.id, a.user_id, u.roll_number, u.name, u.role, u.department, u.year, u.photo_path,
               a.date, a.time, a.status, a.confidence, a.gate_id,
               COALESCE(f.first_seen, a.time) as first_seen,
               COALESCE(f.entry_gate, a.gate_id) as entry_gate,
               COALESCE(f.last_seen, a.time) as last_seen,
               COALESCE(f.last_gate, a.gate_id) as last_gate,
               COALESCE(f.seen_count, 1) as seen_count,
               COALESCE(f.trail_json, '[]') as trail_json
        FROM attendance_events a
        JOIN users u ON a.user_id = u.id
        LEFT JOIN person_footprints f ON f.user_id = a.user_id AND f.date = a.date
        WHERE a.date = ?
    """
    params = [target_date]
    if department and department != "All":
        query += " AND u.department = ?"
        params.append(department)
    if role and role != "All":
        query += " AND u.role = ?"
        params.append(role)

    query += " ORDER BY a.id DESC;"

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(query, params) as cursor:
            rows = await cursor.fetchall()
            results = []
            for row in rows:
                item = dict(row)
                try:
                    item["trail"] = json.loads(item.get("trail_json", "[]"))
                except Exception:
                    item["trail"] = []
                results.append(item)
            return results

@router.get("/footprints")
async def list_footprints(date: Optional[str] = None):
    """Returns list of daily person footprints (when entered, where entered, last seen, and movement trail)."""
    target_date = date or datetime.now().strftime("%Y-%m-%d")
    query = """
        SELECT f.id, f.user_id, u.name, u.roll_number, u.role, u.department, u.photo_path,
               f.date, f.first_seen, f.last_seen, f.entry_gate, f.last_gate, f.seen_count, f.trail_json,
               COALESCE(a.status, 'Present') as status,
               COALESCE(a.confidence, 0.95) as confidence
        FROM person_footprints f
        JOIN users u ON f.user_id = u.id
        LEFT JOIN attendance_events a ON a.user_id = f.user_id AND a.date = f.date
        WHERE f.date = ?
        ORDER BY f.last_seen DESC;
    """
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(query, (target_date,)) as cursor:
            rows = await cursor.fetchall()
            results = []
            for row in rows:
                item = dict(row)
                try:
                    item["trail"] = json.loads(item.get("trail_json", "[]"))
                except Exception:
                    item["trail"] = []
                results.append(item)
            return results

@router.get("/absent")
async def list_absent_records(date: Optional[str] = None, department: Optional[str] = None, role: Optional[str] = None):
    """Returns list of enrolled students/faculty who have NOT checked in on target date."""
    target_date = date or datetime.now().strftime("%Y-%m-%d")
    
    query = """
        SELECT u.id, u.roll_number, u.name, u.role, u.department, u.year, u.email, u.photo_path
        FROM users u
        WHERE u.is_active = 1
        AND u.id NOT IN (
            SELECT user_id FROM attendance_events WHERE date = ?
        )
    """
    params = [target_date]
    if department and department != "All":
        query += " AND u.department = ?"
        params.append(department)
    if role and role != "All":
        query += " AND u.role = ?"
        params.append(role)

    query += " ORDER BY u.department, u.roll_number;"

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(query, params) as cursor:
            rows = await cursor.fetchall()
            return [dict(row) for row in rows]

@router.get("/export/excel")
async def export_attendance_excel(date: Optional[str] = None):
    """Exports structured Excel workbook containing Present and Absent sheets."""
    target_date = date or datetime.now().strftime("%Y-%m-%d")

    wb = openpyxl.Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    late_fill = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")
    present_fill = PatternFill(start_color="EAFAF1", end_color="EAFAF1", fill_type="solid")
    border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    # Sheet 1: Present Records
    ws_present = wb.active
    ws_present.title = "Present Records"
    headers_present = ["Roll Number", "Name", "Role", "Department", "Year", "Date", "Check-in Time", "Status", "Confidence", "Gate"]
    ws_present.append(headers_present)

    for col_idx in range(1, len(headers_present) + 1):
        cell = ws_present.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("""
            SELECT u.roll_number, u.name, u.role, u.department, u.year, a.date, a.time, a.status, a.confidence, a.gate_id
            FROM attendance_events a
            JOIN users u ON a.user_id = u.id
            WHERE a.date = ?
            ORDER BY a.time ASC;
        """, (target_date,)) as cursor:
            present_rows = await cursor.fetchall()
            for r_idx, row in enumerate(present_rows, start=2):
                ws_present.append([
                    row["roll_number"], row["name"], row["role"].capitalize(), row["department"],
                    row["year"], row["date"], row["time"], row["status"], f"{row['confidence']:.2f}", row["gate_id"]
                ])
                status_cell = ws_present.cell(row=r_idx, column=8)
                status_cell.fill = late_fill if row["status"] == "Late" else present_fill

        # Sheet 2: Absent List
        ws_absent = wb.create_sheet(title="Absent Students & Faculty")
        headers_absent = ["Roll Number", "Name", "Role", "Department", "Year", "Email", "Date", "Status"]
        ws_absent.append(headers_absent)
        for col_idx in range(1, len(headers_absent) + 1):
            cell = ws_absent.cell(row=1, column=col_idx)
            cell.fill = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        async with db.execute("""
            SELECT u.roll_number, u.name, u.role, u.department, u.year, u.email
            FROM users u
            WHERE u.is_active = 1
            AND u.id NOT IN (SELECT user_id FROM attendance_events WHERE date = ?)
            ORDER BY u.department, u.roll_number;
        """, (target_date,)) as cursor:
            absent_rows = await cursor.fetchall()
            for row in absent_rows:
                ws_absent.append([
                    row["roll_number"], row["name"], row["role"].capitalize(), row["department"],
                    row["year"], row["email"] or "N/A", target_date, "ABSENT"
                ])

    # Auto-adjust column widths
    for sheet in [ws_present, ws_absent]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Campus_Attendance_{target_date}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
